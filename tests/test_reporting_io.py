"""Regression tests for atomic workbook replacement."""

from __future__ import annotations

from openpyxl import load_workbook
import pandas as pd
import pytest

from Greenhub.reporting_io import atomic_excel_writer, atomic_save_workbook


def test_atomic_excel_writer_and_openpyxl_save(tmp_path):
    path = tmp_path / "report.xlsx"
    with atomic_excel_writer(path) as writer:
        pd.DataFrame({"value": [1.0, 2.0]}).to_excel(
            writer,
            sheet_name="Data",
            index=False,
        )

    workbook = load_workbook(path)
    assert workbook["Data"]["A2"].value == 1.0
    workbook["Data"]["B1"] = "verified"
    atomic_save_workbook(workbook, path)
    workbook.close()

    verified = load_workbook(path, read_only=True, data_only=True)
    assert verified["Data"]["B1"].value == "verified"
    verified.close()
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


def test_atomic_excel_writer_preserves_previous_file_on_failure(tmp_path):
    path = tmp_path / "report.xlsx"
    with atomic_excel_writer(path) as writer:
        pd.DataFrame({"value": ["original"]}).to_excel(writer, index=False)
    original = path.read_bytes()

    with pytest.raises(RuntimeError, match="deliberate"):
        with atomic_excel_writer(path) as writer:
            pd.DataFrame({"value": ["replacement"]}).to_excel(writer, index=False)
            raise RuntimeError("deliberate test failure")

    assert path.read_bytes() == original
    assert not list(tmp_path.glob(".*.tmp.xlsx"))
