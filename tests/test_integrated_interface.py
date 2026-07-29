"""Regression tests for the compatibility-first integrated alpha interface."""

from __future__ import annotations

import csv

from openpyxl import load_workbook

from Greenhub.integrated_reporting import save_integrated_reports
from Greenhub.integration import create_integrated_greenhub_system
from Greenhub.model_completeness import VALID_STATES, write_model_completeness_reports
from Greenhub.systems import create_greenhub_system


def test_integrated_interface_and_reports(tmp_path):
    model = create_integrated_greenhub_system().simulate()

    assert model.legacy_model.system is model.system
    assert model.legacy_model.flowsheet is model.flowsheet
    assert model.scenario_configuration.scenario_chain == "F2/PR3/AD1/BG2/NR4"
    assert model.validation_report is not None
    assert abs(model.validation_report.overall_error_kg_hr) < 1e-6

    process_unit_ids = {
        unit_id
        for area in model.process_areas.values()
        for unit_id in area.units
    }
    assert process_unit_ids == set(model.units)
    assert model.get_process_area("protein_recovery").units
    assert model.get_process_area("shared_utilities_and_supporting_operations").units == {}

    completeness_paths = write_model_completeness_reports(
        model,
        csv_path=tmp_path / "model_completeness.csv",
        markdown_path=tmp_path / "model_completeness.md",
    )
    reporting_paths = save_integrated_reports(
        model,
        output_dir=tmp_path,
        completeness=model.model_completeness,
    )
    completeness_ids = {record.unit_id for record in model.model_completeness}
    assert len(model.model_completeness) > len(model.units)
    assert set(model.units).issubset(completeness_ids)
    assert {
        "T101",
        "S101",
        "M620",
        "M628",
        "S622",
        "M630",
        "S631",
        "M621",
        "A631",
        "A641",
        "R712",
        "A731",
        "E741",
        "M741",
    }.issubset(completeness_ids)
    assert all(record.current_maturity_level in {f"E{i}" for i in range(6)} for record in model.model_completeness)
    assert not any(record.current_maturity_level == "E5" for record in model.model_completeness)
    assert all(
        getattr(record, name) in VALID_STATES
        for record in model.model_completeness
        for name in record.__dataclass_fields__
        if name.endswith("_present")
    )

    with completeness_paths["model_completeness_csv"].open(encoding="utf-8") as file:
        rows = list(csv.DictReader(file))
    assert len(rows) == len(model.model_completeness)
    assert rows[0]["unit_id"]
    assert "purchase_cost_present" in rows[0]
    assert completeness_paths["model_completeness_markdown"].is_file()
    assert reporting_paths["integrated_process_area_summary"].is_file()

    workbook = load_workbook(
        reporting_paths["integrated_greenhub_report"],
        read_only=True,
        data_only=True,
    )
    assert {
        "Run Metadata",
        "Process Areas",
        "Integrated Balances",
        "Products and Residuals",
        "Utilities",
        "Equipment Completeness",
        "Limitations",
    }.issubset(workbook.sheetnames)
    workbook.close()


def test_repeated_integrated_builds_and_legacy_factory_remain_isolated():
    first = create_integrated_greenhub_system()
    first_flowsheet = first.flowsheet
    first_system = first.system
    first_ids = tuple(first.units)

    second = create_integrated_greenhub_system()
    assert second.flowsheet is not first_flowsheet
    assert second.system is not first_system
    assert tuple(second.units) == first_ids

    historical_foundation = create_greenhub_system()
    assert set(historical_foundation.units) == {"T101", "S101"}
    assert historical_foundation.system is not second.system
