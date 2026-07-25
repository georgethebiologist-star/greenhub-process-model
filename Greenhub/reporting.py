"""Phase 2 CSV and Excel reporting utilities."""

from __future__ import annotations

from pathlib import Path

import biosteam as bst
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from Greenhub.analytical import calculate_elemental_flows
from Greenhub.chemicals import chemicals
from Greenhub.component_metadata import metadata_table
from Greenhub.parameters import parameter_table
from Greenhub.reactions import reaction_library_table
from Greenhub.settings import RESULTS_DIR, ensure_results_dir
from Greenhub.stream_templates import stream_template_table
from Greenhub.validation import separator_partitioning_rows


_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_ASSUMED_FILL = PatternFill("solid", fgColor="FFF2CC")
_WARNING_FILL = PatternFill("solid", fgColor="FCE4D6")


def _format_workbook(workbook) -> None:
    """Apply a compact, readable research-report format to every data sheet."""

    for worksheet in workbook.worksheets:
        if worksheet.title == "Flowsheet":
            worksheet.sheet_view.showGridLines = False
            continue

        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            if cell.value is not None:
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 30

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                value = str(cell.value).strip().lower() if cell.value is not None else ""
                if value in {"assumed", "e", "warning", "large_gap"}:
                    cell.fill = _ASSUMED_FILL if value in {"assumed", "e"} else _WARNING_FILL

        for column_index in range(1, worksheet.max_column + 1):
            values = [
                str(worksheet.cell(row=row, column=column_index).value or "")
                for row in range(1, min(worksheet.max_row, 250) + 1)
            ]
            width = max((len(line) for value in values for line in value.splitlines()), default=0)
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(width + 2, 12), 48)

        if worksheet.title == "Validation":
            worksheet.column_dimensions["A"].width = 100
            worksheet.row_dimensions[2].height = 210


def _write_unit_tables(writer: pd.ExcelWriter, sheet_name: str, tables) -> None:
    start_row = 0
    for table in tables:
        if table is None or table.empty:
            continue
        table.to_excel(writer, sheet_name=sheet_name, startrow=start_row)
        start_row += len(table.index) + 3


def _indicator_rows(indicators: dict[str, float | None]) -> list[dict[str, object]]:
    rows = []
    for name, value in indicators.items():
        unit = "-"
        for suffix, resolved in (
            ("_kg_hr", "kg/hr"),
            ("_mg_L", "mg/L"),
            ("_g_L", "g/L"),
            ("_m3_hr", "m3/hr"),
            ("_fraction", "fraction"),
            ("_ratio", "ratio"),
        ):
            if name.endswith(suffix):
                unit = resolved
                break
        rows.append({"indicator": name, "value": value, "unit": unit})
    return rows


def _component_rows(model) -> list[dict[str, object]]:
    flows = model.characterization.component_flows_kg_hr
    return [
        {
            "component_id": component_id,
            "flow_kg_hr": float(flows.get(component_id, 0.0)),
            "mass_fraction": float(flows.get(component_id, 0.0)) / model.streams["feed"].F_mass,
            "input_status": "allocated" if component_id in flows else "not_supplied",
        }
        for component_id in chemicals.IDs
    ]


def _element_balance_rows(model, validation) -> list[dict[str, object]]:
    feed_flows = {ID: float(model.streams["feed"].imass[ID]) for ID in chemicals.IDs}
    outlet_flows = {
        ID: float(model.streams["clarified_liquid"].imass[ID] + model.streams["wet_solids"].imass[ID])
        for ID in chemicals.IDs
    }
    inlet = calculate_elemental_flows(feed_flows)
    outlet = calculate_elemental_flows(outlet_flows)
    return [
        {
            "element": element,
            "inlet_kg_hr": inlet[element],
            "outlet_kg_hr": outlet[element],
            "residual_kg_hr": validation.element_errors_kg_hr[element],
        }
        for element in inlet
    ]


def save_phase2_reports(
    model,
    validation,
    *,
    flowsheet_path: Path,
    results_dir: Path = RESULTS_DIR,
) -> dict[str, Path]:
    """Write all required characterization, partitioning, and system reports."""

    ensure_results_dir()
    results_dir.mkdir(parents=True, exist_ok=True)
    tables = {
        "feed_component_flows.csv": pd.DataFrame(_component_rows(model)),
        "input_reconciliation.csv": pd.DataFrame([row.as_dict() for row in model.characterization.reconciliation]),
        "analytical_indicators.csv": pd.DataFrame(_indicator_rows(model.characterization.indicators)),
        "elemental_balance.csv": pd.DataFrame(_element_balance_rows(model, validation)),
        "separator_partitioning.csv": pd.DataFrame(separator_partitioning_rows(model)),
        "data_quality_report.csv": pd.DataFrame(model.characterization.data_quality_rows),
        "reaction_library.csv": pd.DataFrame(reaction_library_table()),
    }
    output_paths: dict[str, Path] = {"flowsheet": flowsheet_path}
    for filename, table in tables.items():
        path = results_dir / filename
        table.to_csv(path, index=False)
        output_paths[filename] = path

    characterization_path = results_dir / "feed_characterization.xlsx"
    with pd.ExcelWriter(characterization_path, engine="openpyxl") as writer:
        tables["feed_component_flows.csv"].to_excel(writer, sheet_name="Component flows", index=False)
        pd.DataFrame(model.characterization.allocation_rows).to_excel(writer, sheet_name="Input allocation", index=False)
        tables["analytical_indicators.csv"].to_excel(writer, sheet_name="Indicators", index=False)
        tables["input_reconciliation.csv"].to_excel(writer, sheet_name="Reconciliation", index=False)
        tables["data_quality_report.csv"].to_excel(writer, sheet_name="Data quality", index=False)
        pd.DataFrame(metadata_table()).to_excel(writer, sheet_name="Component metadata", index=False)
        pd.DataFrame(parameter_table()).to_excel(writer, sheet_name="Parameters", index=False)
        pd.DataFrame(stream_template_table()).to_excel(writer, sheet_name="Stream templates", index=False)
        tables["reaction_library.csv"].to_excel(writer, sheet_name="Reaction library", index=False)
        _format_workbook(writer.book)
    output_paths["feed_characterization"] = characterization_path

    system_report_path = results_dir / "phase2_system_report.xlsx"
    with pd.ExcelWriter(system_report_path, engine="openpyxl") as writer:
        model.system.results().to_excel(writer, sheet_name="System results")
        bst.report.stream_table(
            model.system.streams,
            flow="kg/hr",
            percent=False,
            T="K",
            P="Pa",
        ).to_excel(writer, sheet_name="Stream table")
        _write_unit_tables(
            writer,
            "Unit results",
            bst.report.unit_result_tables(
                model.system.units,
                include_utilities=True,
                include_total_cost=True,
                include_installed_cost=True,
            ),
        )
        tables["separator_partitioning.csv"].to_excel(writer, sheet_name="Separator", index=False)
        tables["elemental_balance.csv"].to_excel(writer, sheet_name="Element balance", index=False)
        tables["analytical_indicators.csv"].to_excel(writer, sheet_name="Indicators", index=False)
        tables["input_reconciliation.csv"].to_excel(writer, sheet_name="Reconciliation", index=False)
        pd.DataFrame(
            [{"validation_report": validation.report}]
        ).to_excel(writer, sheet_name="Validation", index=False)
        if flowsheet_path.exists():
            from openpyxl.drawing.image import Image as XLImage

            worksheet = writer.book.create_sheet("Flowsheet")
            worksheet.add_image(XLImage(str(flowsheet_path)), "A1")
        _format_workbook(writer.book)
    output_paths["system_report"] = system_report_path
    return output_paths
