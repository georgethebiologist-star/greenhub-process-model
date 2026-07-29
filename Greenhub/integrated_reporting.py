"""Process-area and whole-plant reporting for the integrated GreenHub facade."""

from __future__ import annotations

from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
import subprocess

import biosteam as bst
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import thermosteam as tmo

from Greenhub.model_completeness import UnitCompletenessRecord, build_model_completeness
from Greenhub.reporting_io import atomic_excel_writer, atomic_save_workbook
from Greenhub.settings import RESULTS_DIR

HEADER_FILL = PatternFill("solid", fgColor="244C3D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_BORDER = Border(bottom=Side(style="medium", color="9BB8A5"))
ALTERNATE_FILL = PatternFill("solid", fgColor="EFF5F0")


def _format_integrated_workbook(path: Path) -> None:
    """Apply compact, readable formatting without changing report values."""

    workbook = load_workbook(path)
    narrative_widths = {
        "Run Metadata": {"A": 28, "B": 85},
        "Process Areas": {"A": 48, "B": 18, "C": 45, "D": 35, "E": 45, "F": 36},
        "Dataset Provenance": {"A": 40, "B": 100},
        "Assumptions": {"A": 120},
        "Warnings": {"A": 120},
        "Limitations": {"A": 18, "B": 12, "C": 85, "D": 85},
    }
    area_sheets = {
        "Feedstock",
        "Preparation",
        "Protein Recovery",
        "Anaerobic Digestion",
        "Biogas Upgrading",
        "Nutrient Recovery",
        "Shared Utilities",
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 34
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = HEADER_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row_index in range(2, sheet.max_row + 1):
            if row_index % 2 == 0:
                for cell in sheet[row_index]:
                    cell.fill = ALTERNATE_FILL
            for cell in sheet[row_index]:
                cell.alignment = Alignment(vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"

        widths = narrative_widths.get(sheet.title, {})
        if sheet.title in area_sheets:
            widths = {"A": 18, "B": 30, "C": 45, "D": 20, "E": 90}
        elif sheet.title == "Integrated Balances":
            widths = {"A": 24, "B": 20, "C": 20, "D": 20, "E": 14}
        elif sheet.title == "Products and Residuals":
            widths = {"A": 32, "B": 32, "C": 20, "D": 18, "E": 18, "F": 16, "G": 16, "H": 12, "I": 20}
        elif sheet.title == "Utilities":
            widths = {"A": 12, "B": 48, "C": 25, "D": 25, "E": 20, "F": 20}
        elif sheet.title == "Equipment Completeness":
            widths = {
                get_column_letter(index): (32 if index in {3, 28, 29, 30} else 20)
                for index in range(1, sheet.max_column + 1)
            }
        if not widths:
            for column_index in range(1, sheet.max_column + 1):
                letter = get_column_letter(column_index)
                longest = max(
                    (len(str(cell.value)) for cell in sheet[letter] if cell.value is not None),
                    default=10,
                )
                widths[letter] = min(max(longest + 2, 12), 45)
        for letter, width in widths.items():
            sheet.column_dimensions[letter].width = width

        wrap_columns = set()
        if sheet.title in {"Assumptions", "Warnings"}:
            wrap_columns = {"A"}
        elif sheet.title in {"Dataset Provenance", "Limitations"}:
            wrap_columns = set(widths)
        elif sheet.title in area_sheets:
            wrap_columns = {"C", "E"}
        for letter in wrap_columns:
            for cell in sheet[letter][1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if wrap_columns:
            for row_index in range(2, sheet.max_row + 1):
                sheet.row_dimensions[row_index].height = 32
    atomic_save_workbook(workbook, path)


def _git_commit() -> str:
    try:
        completed = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            check=True,
            capture_output=True,
            text=True,
        )
    except (OSError, subprocess.CalledProcessError):
        return "unavailable"
    return completed.stdout.strip()


def _validation_status(validation: object | None) -> str:
    if validation is None:
        return "not run or not separately defined"
    return "passed; validator raises on failure"


def _process_area_rows(model) -> list[dict[str, object]]:
    rows = []
    for view in model.process_areas.values():
        rows.append(
            {
                "process_area": view.name,
                "active_unit_count": len(view.units),
                "active_units": ", ".join(view.units) or "none",
                "inlet_streams": ", ".join(view.inlet_streams) or "none",
                "outlet_streams": ", ".join(view.outlet_streams) or "none",
                "validation": _validation_status(view.validation),
                "warning_count": len(view.warnings),
                "assumption_count": len(view.assumptions),
            }
        )
    return rows


def _balance_rows(model) -> list[dict[str, object]]:
    validation = model.validation_report
    if validation is None:
        return [{"balance": "integrated validation", "in_kg_hr": None, "out_kg_hr": None, "error_kg_hr": None, "status": "not run"}]
    rows = [
        {
            "balance": "overall mass",
            "in_kg_hr": getattr(validation, "boundary_in_kg_hr", None),
            "out_kg_hr": getattr(validation, "boundary_out_kg_hr", None),
            "error_kg_hr": getattr(validation, "overall_error_kg_hr", None),
            "status": "passed",
        },
        {
            "balance": "water component",
            "in_kg_hr": None,
            "out_kg_hr": None,
            "error_kg_hr": getattr(validation, "water_component_error_kg_hr", None),
            "status": "passed",
        },
    ]
    element_in = getattr(validation, "element_in_kg_hr", {})
    element_out = getattr(validation, "element_out_kg_hr", {})
    element_errors = getattr(validation, "element_errors_kg_hr", {})
    for element, error in element_errors.items():
        rows.append(
            {
                "balance": element,
                "in_kg_hr": element_in.get(element),
                "out_kg_hr": element_out.get(element),
                "error_kg_hr": error,
                "status": "passed",
            }
        )
    return rows


def _product_rows(model) -> list[dict[str, object]]:
    names = (
        "dry_protein_product",
        "biomethane",
        "reject_gas",
        "digestate",
        "nutrient_rich_solids",
        "phosphorus_product",
        "ammonium_product",
        "potassium_concentrate",
        "treated_water",
        "residual_liquor",
        "phase7_wastewater",
        "phase7_offgas",
        "uncaptured_biogas",
    )
    rows = []
    for name in names:
        stream = model.streams.get(name)
        if stream is None:
            continue
        water = float(stream.imass["Water"])
        rows.append(
            {
                "stream_name": name,
                "stream_id": stream.ID,
                "mass_flow_kg_hr": float(stream.F_mass),
                "water_kg_hr": water,
                "dry_mass_kg_hr": float(stream.F_mass) - water,
                "temperature_K": float(stream.T),
                "pressure_Pa": float(stream.P),
                "phase": stream.phase,
                "eligibility_status": "not assessed",
            }
        )
    return rows


def _utility_rows(model) -> list[dict[str, object]]:
    area_by_unit = {
        unit_id: view.name
        for view in model.process_areas.values()
        for unit_id in view.units
    }
    rows = []
    for unit_id, unit in model.units.items():
        heating = 0.0
        cooling = 0.0
        for utility in unit.heat_utilities:
            duty_kw = float(utility.duty or 0.0) / 3600.0
            if duty_kw >= 0.0:
                heating += duty_kw
            else:
                cooling += abs(duty_kw)
        rows.append(
            {
                "unit_id": unit_id,
                "process_area": area_by_unit.get(unit_id, "Shared Utilities and Supporting Operations"),
                "electricity_consumption_kW": max(float(unit.power_utility.rate), 0.0),
                "electricity_production_kW": max(-float(unit.power_utility.rate), 0.0),
                "heating_duty_kW": heating,
                "cooling_duty_kW": cooling,
            }
        )
    return rows


def _metadata_rows(model) -> list[dict[str, object]]:
    scenario = model.scenario_configuration
    specification = model.feed_specification
    return [
        {"field": "generated_utc", "value": datetime.now(timezone.utc).isoformat()},
        {"field": "git_commit", "value": _git_commit()},
        {"field": "BioSTEAM_version", "value": bst.__version__},
        {"field": "ThermoSTEAM_version", "value": tmo.__version__},
        {"field": "system_id", "value": model.system.ID},
        {"field": "flowsheet_id", "value": model.flowsheet.ID},
        {"field": "input_workbook", "value": model.dataset_provenance["master_input_workbook"]},
        {"field": "scenario_chain", "value": scenario.scenario_chain},
        {"field": "process_order", "value": scenario.process_order},
        {"field": "selected_month", "value": scenario.selected_month},
        {"field": "sample_id", "value": specification.sample_id},
        {"field": "monthly_feed_mass_kg", "value": specification.monthly_total_mass_kg},
        {"field": "average_feed_flow_kg_hr", "value": specification.total_mass_flow_kg_hr},
        {"field": "structural_test_mode", "value": scenario.structural_test_mode},
        {"field": "input_basis", "value": scenario.input_basis},
    ]


def _area_detail_rows(view, completeness_by_id) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for unit_id, unit in view.units.items():
        record = completeness_by_id[unit_id]
        rows.append(
            {
                "record_type": "unit",
                "name": unit_id,
                "value": type(unit).__name__,
                "status": record.current_maturity_level,
                "detail": record.missing_items,
            }
        )
    for name, stream in view.inlet_streams.items():
        rows.append(
            {
                "record_type": "inlet stream",
                "name": name,
                "value": float(stream.F_mass),
                "status": "kg/hr",
                "detail": stream.ID,
            }
        )
    for name, stream in view.outlet_streams.items():
        rows.append(
            {
                "record_type": "outlet stream",
                "name": name,
                "value": float(stream.F_mass),
                "status": "kg/hr",
                "detail": stream.ID,
            }
        )
    for index, warning in enumerate(view.warnings, start=1):
        rows.append({"record_type": "warning", "name": index, "value": warning, "status": "open", "detail": ""})
    for index, assumption in enumerate(view.assumptions, start=1):
        rows.append({"record_type": "assumption", "name": index, "value": assumption, "status": "documented", "detail": ""})
    if not rows:
        rows.append(
            {
                "record_type": "status",
                "name": view.name,
                "value": "No active units or streams in selected configuration",
                "status": "not yet implemented",
                "detail": "",
            }
        )
    return rows


def save_integrated_reports(
    model,
    *,
    output_dir: Path = RESULTS_DIR,
    completeness: tuple[UnitCompletenessRecord, ...] | None = None,
) -> dict[str, Path]:
    """Save an additive, process-area-named CSV and Excel workbook."""

    output_dir.mkdir(parents=True, exist_ok=True)
    completeness = completeness or build_model_completeness(model)
    completeness_by_id = {record.unit_id: record for record in completeness}
    summary_path = output_dir / "integrated_process_area_summary.csv"
    workbook_path = output_dir / "integrated_greenhub_report.xlsx"

    process_area_frame = pd.DataFrame(_process_area_rows(model))
    process_area_frame.to_csv(summary_path, index=False)
    with atomic_excel_writer(workbook_path) as writer:
        pd.DataFrame(_metadata_rows(model)).to_excel(writer, sheet_name="Run Metadata", index=False)
        process_area_frame.to_excel(writer, sheet_name="Process Areas", index=False)
        pd.DataFrame(_balance_rows(model)).to_excel(writer, sheet_name="Integrated Balances", index=False)
        pd.DataFrame(_product_rows(model)).to_excel(writer, sheet_name="Products and Residuals", index=False)
        pd.DataFrame(_utility_rows(model)).to_excel(writer, sheet_name="Utilities", index=False)
        pd.DataFrame(asdict(record) for record in completeness).to_excel(
            writer,
            sheet_name="Equipment Completeness",
            index=False,
        )
        pd.DataFrame(
            ({"dataset": key, "provenance_or_status": value} for key, value in model.dataset_provenance.items())
        ).to_excel(writer, sheet_name="Dataset Provenance", index=False)
        pd.DataFrame(
            ({"assumption": value} for value in model.assumptions)
        ).to_excel(writer, sheet_name="Assumptions", index=False)
        pd.DataFrame(
            ({"warning": value} for value in model.warnings)
        ).to_excel(writer, sheet_name="Warnings", index=False)
        pd.DataFrame(
            (
                {
                    "limitation_id": getattr(item, "limitation_id", ""),
                    "priority": getattr(item, "priority", ""),
                    "description": getattr(item, "description", str(item)),
                    "data_needed": getattr(item, "data_needed", ""),
                }
                for item in model.limitations
            )
        ).to_excel(writer, sheet_name="Limitations", index=False)
        sheet_names = {
            "feedstock_and_characterization": "Feedstock",
            "feed_preparation_and_primary_separation": "Preparation",
            "protein_recovery": "Protein Recovery",
            "anaerobic_digestion": "Anaerobic Digestion",
            "biogas_conditioning_and_upgrading": "Biogas Upgrading",
            "digestate_treatment_and_nutrient_recovery": "Nutrient Recovery",
            "shared_utilities_and_supporting_operations": "Shared Utilities",
        }
        for key, view in model.process_areas.items():
            pd.DataFrame(_area_detail_rows(view, completeness_by_id)).to_excel(
                writer,
                sheet_name=sheet_names[key],
                index=False,
            )

    _format_integrated_workbook(workbook_path)

    paths = {
        "integrated_process_area_summary": summary_path,
        "integrated_greenhub_report": workbook_path,
    }
    model.integrated_report_paths.update(paths)
    return paths
