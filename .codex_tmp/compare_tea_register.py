from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path.cwd()
WORKBOOK = Path(r"C:\Users\katrb\Downloads\GreenHub_TEA_Compact_Cost_Register (1).xlsx")
CATALOG = ROOT / "Greenhub" / "results" / "economic_inventory_catalog.csv"


def norm(text: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def tokens(text: object) -> set[str]:
    stop = {
        "and", "or", "the", "with", "for", "unit", "package", "selected",
        "cost", "costs", "equipment", "system", "plant", "greenhub",
    }
    return {t for t in norm(text).split() if len(t) >= 4 and t not in stop}


wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
ws = wb["Cost Inventory"]
header = [cell for cell in next(ws.iter_rows(min_row=4, max_row=4, values_only=True))]
rows = []
for row in ws.iter_rows(min_row=5, values_only=True):
    if not any(row):
        continue
    rows.append(dict(zip(header, row)))

catalog_rows = []
with CATALOG.open("r", encoding="utf-8-sig", newline="") as handle:
    reader = csv.DictReader(handle)
    for row in reader:
        hay = " ".join(
            row.get(k, "") for k in (
                "process_area", "process_step", "item_category", "item_name",
                "parameter_id", "python_variable", "role_or_description", "notes",
            )
        )
        catalog_rows.append((row, tokens(hay)))

matches = []
for row in rows:
    needle_text = " ".join(str(row.get(k) or "") for k in (
        "process_area", "process_step", "equipment_or_cost_item", "item_category"
    ))
    needle = tokens(needle_text)
    scored = []
    for cat_row, cat_tokens in catalog_rows:
        score = len(needle & cat_tokens)
        if score:
            scored.append((score, cat_row))
    scored.sort(key=lambda pair: pair[0], reverse=True)
    best = scored[:5]
    matches.append({
        "process_area": row.get("process_area"),
        "process_step": row.get("process_step"),
        "item": row.get("equipment_or_cost_item"),
        "category": row.get("item_category"),
        "evidence": row.get("evidence_status"),
        "confidence": row.get("confidence"),
        "value": row.get("converted_value_eur"),
        "unit": row.get("converted_unit"),
        "scaling": row.get("scaling_exponent"),
        "match_count": len(scored),
        "top_match": best[0][1].get("item_name") if best else "",
        "top_match_file": best[0][1].get("source_file") if best else "",
        "top_match_readiness": best[0][1].get("tea_readiness") if best else "",
        "top_score": best[0][0] if best else 0,
    })

print("cost_inventory_rows", len(rows))
print("category_counts", dict(Counter(str(r.get("item_category")) for r in rows)))
print("confidence_counts", dict(Counter(str(r.get("confidence")) for r in rows)))
print("evidence_counts", dict(Counter(str(r.get("evidence_status")) for r in rows)))
print("matched_rows_score_gt0", sum(1 for r in matches if r["top_score"]))
print("strong_matches_score_ge3", sum(1 for r in matches if r["top_score"] >= 3))
print("\nITEMS")
for m in matches:
    print(
        f"- {m['process_area']} | {m['process_step']} | {m['item']} | "
        f"{m['category']} | EUR={m['value']} {m['unit']} | scaling={m['scaling']} | "
        f"match_score={m['top_score']} | match={m['top_match']} | file={m['top_match_file']}"
    )
