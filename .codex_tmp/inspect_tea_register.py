from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(r"C:\Users\katrb\Downloads\GreenHub_TEA_Compact_Cost_Register (1).xlsx")

wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\nSHEET\t{ws.title}\trows={ws.max_row}\tcols={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True):
        print("\t".join("" if value is None else str(value) for value in row))
